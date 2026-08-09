"""Tests for the configurable export system (app.services.export_service +
the /api/v1/export contract routes in export_report.py / bom_enterprise.py).

Covers the contract's PROVE IT list: column selection+order in CSV, unknown
column -> 400 (not silently dropped), indented BOM export levels + quantity
roll-up, xlsx output is a real openpyxl-readable workbook, and tenant
isolation.
"""

import csv
import io

import pytest

from app.models.bom import BOM, BOMItem
from app.models.part import Part
from app.models.tenant import Tenant


async def _make_part(db_session, tenant_id, pn, name="Part", category="Electrical", cost=10.0):
    part = Part(pn=pn, name=name, category=category, cost=cost, tenantId=tenant_id)
    db_session.add(part)
    await db_session.commit()
    await db_session.refresh(part)
    return part


async def _make_bom(db_session, tenant_id, bom_number, name="BOM"):
    bom = BOM(bom_number=bom_number, name=name, tenantId=tenant_id)
    db_session.add(bom)
    await db_session.commit()
    await db_session.refresh(bom)
    return bom


async def _make_item(db_session, tenant_id, bom_id, part_id=None, quantity=1, parent_item_id=None):
    item = BOMItem(
        bom_id=bom_id,
        part_id=part_id,
        quantity=quantity,
        parent_item_id=parent_item_id,
        tenantId=tenant_id,
    )
    db_session.add(item)
    await db_session.commit()
    await db_session.refresh(item)
    return item


def _rows(csv_text: str) -> list[list[str]]:
    return list(csv.reader(io.StringIO(csv_text)))


# ---- GET /export/columns ----


@pytest.mark.asyncio
async def test_export_columns_lists_parts_columns_with_default_flag(client, auth_headers):
    resp = await client.get("/api/v1/export/columns?entity=parts", headers=auth_headers)
    assert resp.status_code == 200
    cols = resp.json()["columns"]
    keys = {c["key"] for c in cols}
    assert "pn" in keys and "cost" in keys
    pn_col = next(c for c in cols if c["key"] == "pn")
    assert pn_col["default"] is True
    assert pn_col["label"]


@pytest.mark.asyncio
async def test_export_columns_unknown_entity_400(client, auth_headers):
    resp = await client.get("/api/v1/export/columns?entity=bogus", headers=auth_headers)
    assert resp.status_code == 400


# ---- CSV column selection + order ----


@pytest.mark.asyncio
async def test_csv_export_honours_requested_column_order(db_session, test_tenant, client, auth_headers):
    await _make_part(db_session, test_tenant.id, pn="PN-A", name="Alpha", cost=5)
    await _make_part(db_session, test_tenant.id, pn="PN-B", name="Beta", cost=7)

    resp = await client.post(
        "/api/v1/export",
        headers=auth_headers,
        json={"entity": "parts", "format": "csv", "columns": ["name", "pn"]},
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    assert "attachment" in resp.headers["content-disposition"]

    rows = _rows(resp.text)
    assert rows[0] == ["Name", "Part Number"]  # requested order, not catalogue order
    assert rows[1] == ["Alpha", "PN-A"]
    assert rows[2] == ["Beta", "PN-B"]


@pytest.mark.asyncio
async def test_unknown_column_is_a_400_not_silently_dropped(db_session, test_tenant, client, auth_headers):
    await _make_part(db_session, test_tenant.id, pn="PN-A", name="Alpha")

    resp = await client.post(
        "/api/v1/export",
        headers=auth_headers,
        json={"entity": "parts", "format": "csv", "columns": ["pn", "not_a_real_column"]},
    )
    assert resp.status_code == 400
    assert "not_a_real_column" in resp.text


@pytest.mark.asyncio
async def test_unknown_entity_and_format_are_400(client, auth_headers):
    resp = await client.post(
        "/api/v1/export", headers=auth_headers, json={"entity": "bogus", "format": "csv"}
    )
    assert resp.status_code == 400

    resp = await client.post(
        "/api/v1/export", headers=auth_headers, json={"entity": "parts", "format": "bogus"}
    )
    assert resp.status_code == 400


# ---- xlsx is a real, readable workbook ----


@pytest.mark.asyncio
async def test_xlsx_export_is_a_real_openpyxl_workbook(db_session, test_tenant, client, auth_headers):
    from openpyxl import load_workbook

    await _make_part(db_session, test_tenant.id, pn="PN-X", name="Xlsx Part", cost=12.5)

    resp = await client.post(
        "/api/v1/export",
        headers=auth_headers,
        json={"entity": "parts", "format": "xlsx", "columns": ["pn", "name", "cost"]},
    )
    assert resp.status_code == 200
    assert "spreadsheet" in resp.headers["content-type"]

    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert [c.value for c in ws[1]] == ["Part Number", "Name", "Cost"]
    assert list(ws[2]) [0].value == "PN-X"


# ---- BOM indented export: levels + quantity roll-up ----


@pytest.mark.asyncio
async def test_bom_export_indented_has_correct_levels_and_qty_rollup(
    db_session, test_tenant, client, auth_headers
):
    tid = test_tenant.id
    parent_part = await _make_part(db_session, tid, pn="PN-PARENT", name="Parent", cost=10)
    child_part = await _make_part(db_session, tid, pn="PN-CHILD", name="Child", cost=2)

    bom = await _make_bom(db_session, tid, bom_number="BOM-EXP-001")
    parent_item = await _make_item(db_session, tid, bom.id, part_id=parent_part.id, quantity=2)
    await _make_item(
        db_session, tid, bom.id, part_id=child_part.id, quantity=3, parent_item_id=parent_item.id
    )

    resp = await client.post(
        "/api/v1/export",
        headers=auth_headers,
        json={
            "entity": "bom",
            "bom_id": bom.id,
            "format": "csv",
            "indented": True,
            "columns": ["level", "pn", "quantity"],
        },
    )
    assert resp.status_code == 200
    rows = _rows(resp.text)
    assert rows[0] == ["Level", "Part Number", "Quantity"]
    # Parent first (correct parent-child order), level 1, qty 2.
    assert rows[1] == ["1", "PN-PARENT", "2.0"]
    # Child: level 2, EFFECTIVE (rolled-up) qty = 3 * 2 = 6, not raw 3.
    assert rows[2] == ["2", "PN-CHILD", "6.0"]


@pytest.mark.asyncio
async def test_bom_export_flat_uses_raw_quantity_not_rollup(db_session, test_tenant, client, auth_headers):
    tid = test_tenant.id
    parent_part = await _make_part(db_session, tid, pn="PN-PARENT2", name="Parent")
    child_part = await _make_part(db_session, tid, pn="PN-CHILD2", name="Child")

    bom = await _make_bom(db_session, tid, bom_number="BOM-EXP-002")
    parent_item = await _make_item(db_session, tid, bom.id, part_id=parent_part.id, quantity=2)
    await _make_item(
        db_session, tid, bom.id, part_id=child_part.id, quantity=3, parent_item_id=parent_item.id
    )

    resp = await client.post(
        "/api/v1/export",
        headers=auth_headers,
        json={
            "entity": "bom",
            "bom_id": bom.id,
            "format": "csv",
            "indented": False,
            "columns": ["pn", "quantity"],
        },
    )
    assert resp.status_code == 200
    rows = _rows(resp.text)
    assert rows[1] == ["PN-PARENT2", "2.0"]
    assert rows[2] == ["PN-CHILD2", "3.0"]  # raw line qty, not 6


@pytest.mark.asyncio
async def test_bom_entity_requires_bom_id(client, auth_headers):
    resp = await client.post(
        "/api/v1/export", headers=auth_headers, json={"entity": "bom", "format": "csv"}
    )
    assert resp.status_code == 400


# ---- Tenant isolation ----


@pytest.mark.asyncio
async def test_export_never_crosses_tenants(db_session, test_tenant, client, auth_headers):
    other = Tenant(id=test_tenant.id + 1, tenant_name="Other Tenant", tenant_code="OTHER2")
    db_session.add(other)
    await db_session.commit()

    await _make_part(db_session, test_tenant.id, pn="PN-MINE", name="Mine")
    # Insert directly (no refresh) — under the active tenant-1 context, a
    # post-insert refresh's SELECT would itself be auto-filtered to tenant 1
    # and find zero rows for this tenant-2 part, raising InvalidRequestError.
    db_session.add(Part(pn="PN-THEIRS", name="Theirs", category="Electrical", tenantId=other.id))
    await db_session.commit()

    resp = await client.post(
        "/api/v1/export",
        headers=auth_headers,
        json={"entity": "parts", "format": "csv", "columns": ["pn"]},
    )
    assert resp.status_code == 200
    assert "PN-MINE" in resp.text
    assert "PN-THEIRS" not in resp.text


# ---- Templates ----


@pytest.mark.asyncio
async def test_template_crud_and_explicit_fields_override_template(
    db_session, test_tenant, client, auth_headers
):
    await _make_part(db_session, test_tenant.id, pn="PN-T", name="Templated", cost=9)

    create_resp = await client.post(
        "/api/v1/export/templates",
        headers=auth_headers,
        json={
            "name": "My Parts Template",
            "entity": "parts",
            "config": {"format": "csv", "columns": ["pn", "name"]},
        },
    )
    assert create_resp.status_code == 201
    tmpl_id = create_resp.json()["id"]

    list_resp = await client.get(
        "/api/v1/export/templates?entity=parts", headers=auth_headers
    )
    assert list_resp.status_code == 200
    assert any(t["id"] == tmpl_id for t in list_resp.json())

    # template_id supplies format+columns; explicit format="json" overrides the
    # template's "csv" (explicit fields always win).
    export_resp = await client.post(
        "/api/v1/export",
        headers=auth_headers,
        json={"entity": "parts", "format": "json", "template_id": tmpl_id},
    )
    assert export_resp.status_code == 200
    assert export_resp.headers["content-type"].startswith("application/json")
    data = export_resp.json()
    assert data == [{"pn": "PN-T", "name": "Templated"}]

    del_resp = await client.delete(
        f"/api/v1/export/templates/{tmpl_id}", headers=auth_headers
    )
    assert del_resp.status_code == 204

    list_resp2 = await client.get(
        "/api/v1/export/templates?entity=parts", headers=auth_headers
    )
    assert all(t["id"] != tmpl_id for t in list_resp2.json())


# ---- Currency: honest error, no fabricated rate ----


@pytest.mark.asyncio
async def test_currency_conversion_without_a_rate_on_file_is_a_400(
    db_session, test_tenant, client, auth_headers
):
    await _make_part(db_session, test_tenant.id, pn="PN-CUR", name="CurPart", cost=100)

    resp = await client.post(
        "/api/v1/export",
        headers=auth_headers,
        json={"entity": "parts", "format": "csv", "currency": "EUR", "columns": ["pn", "cost"]},
    )
    assert resp.status_code == 400


# ---- bom_enterprise.py delegation (no more silently-ignored format) ----


@pytest.mark.asyncio
async def test_bom_enterprise_export_honours_format(db_session, test_tenant, client, auth_headers):
    tid = test_tenant.id
    part = await _make_part(db_session, tid, pn="PN-ENT", name="Enterprise Part", cost=3)
    bom = await _make_bom(db_session, tid, bom_number="BOM-ENT-001")
    await _make_item(db_session, tid, bom.id, part_id=part.id, quantity=4)

    resp = await client.post(f"/api/v1/bom/{bom.id}/export?format=csv", headers=auth_headers)
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    assert "PN-ENT" in resp.text
