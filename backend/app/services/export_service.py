"""Configurable export service — renders parts/vendors/purchase_orders/bom to
csv/xlsx/pdf/json, honouring column selection+order, filters, currency, and
(for BOM) indented multi-level vs flat + include_sub_assemblies.

Single place all export routes (the new /api/v1/export contract, the
back-compat /export/<entity>/<fmt> routes, and bom_enterprise's
POST /{bom_id}/export) funnel through, so column/format/filter rules never
drift between them.
"""

import csv
import io
import json
from datetime import date
from decimal import Decimal
from typing import Any, Optional

from fastapi import HTTPException
from sqlalchemy import select

from app.models.bom import BOMItem
from app.models.export_template import ExportTemplate
from app.models.part import Part
from app.models.po_models import POHeader
from app.models.vendor import Vendor
from app.services.bom_service import (
    _compute_levels_and_effective_qty,
    _drop_excluded_subtrees,
    get_bom_or_404,
)

ENTITIES = ("bom", "parts", "vendors", "purchase_orders")
FORMATS = ("csv", "xlsx", "pdf", "json")

# ---- Column catalogue: (key, label, default) per entity. This is the single
# source of truth for GET /export/columns AND for validating/ordering the
# `columns` param — the UI's picker and the export renderer can never disagree.
COLUMNS: dict[str, list[tuple[str, str, bool]]] = {
    "parts": [
        ("pn", "Part Number", True),
        ("name", "Name", True),
        ("description", "Description", True),
        ("category", "Category", True),
        ("subCategory", "Sub-Category", True),
        ("manufacturer", "Manufacturer", True),
        ("vendor", "Vendor", True),
        ("cost", "Cost", True),
        ("lead", "Lead (days)", True),
        ("origin", "Origin", True),
        ("status", "Status", True),
        ("material", "Material", True),
        ("weight", "Weight (g)", True),
        ("mpn", "MPN", True),
        ("htsCode", "HTS Code", True),
        ("unspscCode", "UNSPSC Code", True),
        ("uom", "UoM", False),
        ("barcode", "Barcode", False),
        ("freight", "Freight", False),
        ("tax", "Tax", False),
        ("landedCost", "Landed Cost", False),
    ],
    "vendors": [
        ("name", "Name", True),
        ("country", "Country", True),
        ("contactEmail", "Email", True),
        ("contactPhone", "Phone", True),
        ("address", "Address", False),
        ("leadTime", "Lead Time (days)", True),
        ("moq", "MOQ", False),
        ("terms", "Terms", False),
        ("reliabilityRating", "Rating", True),
        ("active", "Active", True),
        ("notes", "Notes", False),
    ],
    "purchase_orders": [
        ("poNumber", "PO Number", True),
        ("poDate", "Date", True),
        ("vendorName", "Vendor", True),
        ("project", "Project", True),
        ("poTotal", "Total", True),
        ("status", "Status", True),
        ("currency", "Currency", False),
        ("payment_terms", "Payment Terms", False),
        ("shipping_method", "Shipping Method", False),
        ("line_count", "Line Count", False),
    ],
    "bom": [
        ("level", "Level", True),
        ("pn", "Part Number", True),
        ("name", "Name", True),
        ("reference_designator", "Reference Designator", True),
        ("find_number", "Find #", False),
        ("quantity", "Quantity", True),
        ("unit", "UoM", True),
        ("category", "Category", True),
        ("vendor", "Vendor", False),
        ("manufacturer", "Manufacturer", False),
        ("unit_cost", "Unit Cost", True),
        ("extended_cost", "Extended Cost", True),
        ("notes", "Notes", False),
        ("status", "Part Status", False),
    ],
}

# Money columns get currency-converted when `currency` is requested.
MONEY_COLUMNS: dict[str, set[str]] = {
    "parts": {"cost", "freight", "tax", "landedCost"},
    "vendors": set(),
    "purchase_orders": {"poTotal"},
    "bom": {"unit_cost", "extended_cost"},
}

# Allow-listed filter keys -> real column, per entity. Anything else is a 400.
FILTERS: dict[str, set[str]] = {
    "parts": {"category", "subCategory", "status", "vendor", "manufacturer", "part_kind"},
    "vendors": {"active", "country"},
    "purchase_orders": {"status", "vendorName", "project"},
    "bom": {"category", "status", "vendor", "manufacturer"},  # applied to the joined Part
}

CONTENT_TYPES = {
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
    "json": "application/json",
}


def _bad_request(msg: str) -> HTTPException:
    return HTTPException(status_code=400, detail=msg)


def validate_entity(entity: str) -> None:
    if entity not in ENTITIES:
        raise _bad_request(f"Unknown entity '{entity}'. Must be one of {ENTITIES}.")


def validate_format(fmt: str) -> None:
    if fmt not in FORMATS:
        raise _bad_request(f"Unknown format '{fmt}'. Must be one of {FORMATS}.")


def resolve_columns(entity: str, columns: Optional[list[str]], indented: bool = True) -> list[str]:
    """Validate + resolve the requested column list, or fall back to defaults."""
    spec = COLUMNS[entity]
    valid_keys = {k for k, _, _ in spec}
    if columns is not None:
        unknown = [c for c in columns if c not in valid_keys]
        if unknown:
            raise _bad_request(f"Unknown column(s) for entity '{entity}': {unknown}")
        return list(columns)
    defaults = [k for k, _, default in spec if default]
    if entity == "bom" and not indented:
        defaults = [k for k in defaults if k != "level"]
    return defaults


def validate_filters(entity: str, filters: Optional[dict[str, Any]]) -> dict[str, Any]:
    if not filters:
        return {}
    allowed = FILTERS[entity]
    unknown = [k for k in filters if k not in allowed]
    if unknown:
        raise _bad_request(f"Unknown filter(s) for entity '{entity}': {unknown}")
    return filters


async def _get_conversion_rate(db, tenant_id: Optional[int], currency: Optional[str]) -> float:
    """Rate to multiply stored (assumed USD) values by. None/USD -> 1.0.
    Raises 400 (never fabricates a rate) if no rate is on file."""
    if not currency or currency == "USD":
        return 1.0
    from app.models.enterprise_extensions import ExchangeRate

    stmt = (
        select(ExchangeRate.rate)
        .where(ExchangeRate.from_currency == "USD", ExchangeRate.to_currency == currency)
        .where(ExchangeRate.is_active.is_(True))
    )
    if tenant_id is not None:
        stmt = stmt.where(ExchangeRate.tenantId == tenant_id)
    stmt = stmt.order_by(ExchangeRate.effective_date.desc()).limit(1)
    rate = (await db.execute(stmt)).scalar()
    if rate is None:
        raise _bad_request(f"No exchange rate on file for USD -> {currency}")
    return float(rate)


def _num(v: Any) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, Decimal):
        return float(v)
    return v


async def _rows_parts(db, tenant_id, filters) -> list[dict]:
    stmt = select(Part)
    if tenant_id is not None:
        stmt = stmt.where(Part.tenantId == tenant_id)
    for key, value in filters.items():
        stmt = stmt.where(getattr(Part, key) == value)
    stmt = stmt.order_by(Part.pn)
    parts = (await db.execute(stmt)).scalars().all()
    return [
        {
            "pn": p.pn,
            "name": p.name,
            "description": p.description,
            "category": p.category,
            "subCategory": p.subCategory,
            "manufacturer": p.manufacturer,
            "vendor": p.vendor,
            "cost": _num(p.cost),
            "lead": p.lead,
            "origin": p.origin,
            "status": p.status,
            "material": p.material,
            "weight": p.weight,
            "mpn": p.mpn,
            "htsCode": p.htsCode,
            "unspscCode": p.unspscCode,
            "uom": p.uom,
            "barcode": p.barcode,
            "freight": _num(p.freight),
            "tax": _num(p.tax),
            "landedCost": _num(p.landedCost),
        }
        for p in parts
    ]


async def _rows_vendors(db, tenant_id, filters) -> list[dict]:
    stmt = select(Vendor)
    if tenant_id is not None:
        stmt = stmt.where(Vendor.tenantId == tenant_id)
    for key, value in filters.items():
        stmt = stmt.where(getattr(Vendor, key) == value)
    stmt = stmt.order_by(Vendor.name)
    vendors = (await db.execute(stmt)).scalars().all()
    return [
        {
            "name": v.name,
            "country": v.country,
            "contactEmail": v.contactEmail,
            "contactPhone": v.contactPhone,
            "address": v.address,
            "leadTime": v.leadTime,
            "moq": v.moq,
            "terms": v.terms,
            "reliabilityRating": v.reliabilityRating,
            "active": v.active,
            "notes": v.notes,
        }
        for v in vendors
    ]


async def _rows_purchase_orders(db, tenant_id, filters) -> list[dict]:
    stmt = select(POHeader)
    if tenant_id is not None:
        stmt = stmt.where(POHeader.tenantId == tenant_id)
    for key, value in filters.items():
        stmt = stmt.where(getattr(POHeader, key) == value)
    stmt = stmt.order_by(POHeader.poDate.desc())
    pos = (await db.execute(stmt)).scalars().all()
    return [
        {
            "poNumber": po.poNumber,
            "poDate": po.poDate,
            "vendorName": po.vendorName,
            "project": po.project,
            "poTotal": _num(po.poTotal),
            "status": po.status,
            "currency": po.currency,
            "payment_terms": po.payment_terms,
            "shipping_method": po.shipping_method,
            "line_count": po.line_count,
        }
        for po in pos
    ]


async def _rows_bom(
    db,
    tenant_id,
    bom_id: int,
    filters: dict,
    indented: bool,
    include_sub_assemblies: bool,
) -> list[dict]:
    bom = await get_bom_or_404(db, bom_id)
    if tenant_id is not None and bom.tenantId != tenant_id:
        raise HTTPException(status_code=404, detail="BOM not found")

    stmt = select(BOMItem).where(BOMItem.bom_id == bom_id)
    if tenant_id is not None:
        stmt = stmt.where(BOMItem.tenantId == tenant_id)
    items = _drop_excluded_subtrees((await db.execute(stmt)).scalars().all())

    if not include_sub_assemblies:
        items = [i for i in items if i.parent_item_id is None]

    part_ids = {i.part_id for i in items if i.part_id}
    parts_map: dict[int, Part] = {}
    if part_ids:
        pr_stmt = select(Part).where(Part.id.in_(part_ids))
        if tenant_id is not None:
            pr_stmt = pr_stmt.where(Part.tenantId == tenant_id)
        parts_map = {p.id: p for p in (await db.execute(pr_stmt)).scalars().all()}

    if filters:
        def _matches(item: BOMItem) -> bool:
            part = parts_map.get(item.part_id)
            return all(getattr(part, key, None) == value for key, value in filters.items())

        items = [i for i in items if _matches(i)]

    levels, effective_qty = _compute_levels_and_effective_qty(items)

    # Parent-first DFS order (correct parent-child order), siblings by sort_order/id.
    by_parent: dict[Optional[int], list[BOMItem]] = {}
    item_ids = {i.id for i in items}
    for it in items:
        parent_key = it.parent_item_id if it.parent_item_id in item_ids else None
        by_parent.setdefault(parent_key, []).append(it)
    for group in by_parent.values():
        group.sort(key=lambda i: (i.sort_order or 0, i.id))

    ordered: list[BOMItem] = []

    def walk(parent_key: Optional[int]) -> None:
        for it in by_parent.get(parent_key, []):
            ordered.append(it)
            walk(it.id)

    walk(None)

    rows = []
    for item in ordered:
        part = parts_map.get(item.part_id)
        unit_cost = item.unit_cost_snapshot if item.unit_cost_snapshot is not None else (
            part.cost if part else None
        )
        qty = effective_qty[item.id] if indented else float(item.quantity or 0)
        rows.append(
            {
                "level": levels[item.id],
                "pn": part.pn if part else None,
                "name": part.name if part else None,
                "reference_designator": item.reference_designator,
                "find_number": item.find_number,
                "quantity": qty,
                "unit": item.unit,
                "category": part.category if part else None,
                "vendor": part.vendor if part else None,
                "manufacturer": part.manufacturer if part else None,
                "unit_cost": _num(unit_cost),
                "extended_cost": (
                    round(float(unit_cost or 0) * qty, 4) if unit_cost is not None else None
                ),
                "notes": item.notes,
                "status": part.status if part else None,
            }
        )
    return rows


def _render_csv(columns: list[tuple[str, str]], rows: list[dict]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([label for _, label in columns])
    for row in rows:
        writer.writerow(["" if row.get(k) is None else row.get(k) for k, _ in columns])
    return buf.getvalue().encode("utf-8")


def _render_json(columns: list[tuple[str, str]], rows: list[dict]) -> bytes:
    keys = [k for k, _ in columns]
    data = [{k: row.get(k) for k in keys} for row in rows]
    return json.dumps(data, indent=2, default=str).encode("utf-8")


def _render_xlsx(columns: list[tuple[str, str]], rows: list[dict], sheet_title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31] or "Export"

    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    for col, (_, label) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for r, row in enumerate(rows, 2):
        for c, (key, _) in enumerate(columns, 1):
            ws.cell(row=r, column=c, value=row.get(key))

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _render_pdf(columns: list[tuple[str, str]], rows: list[dict], title: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4))
    styles = getSampleStyleSheet()
    elements = [Paragraph(title, styles["Title"]), Spacer(1, 12)]

    data = [[label for _, label in columns]]
    for row in rows:
        data.append(["" if row.get(k) is None else str(row.get(k)) for k, _ in columns])

    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, 0), 8),
                ("FONTSIZE", (0, 1), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)
    return buf.getvalue()


async def render_export(
    db,
    tenant_id: Optional[int],
    *,
    entity: str,
    format: str,
    bom_id: Optional[int] = None,
    columns: Optional[list[str]] = None,
    indented: bool = True,
    include_sub_assemblies: bool = True,
    filters: Optional[dict[str, Any]] = None,
    currency: Optional[str] = None,
) -> tuple[bytes, str, str]:
    """Returns (file_bytes, content_type, filename)."""
    validate_entity(entity)
    validate_format(format)
    filters = validate_filters(entity, filters)

    if entity == "bom" and bom_id is None:
        raise _bad_request("bom_id is required when entity == 'bom'")

    resolved_columns = resolve_columns(entity, columns, indented=indented)

    if entity == "parts":
        rows = await _rows_parts(db, tenant_id, filters)
    elif entity == "vendors":
        rows = await _rows_vendors(db, tenant_id, filters)
    elif entity == "purchase_orders":
        rows = await _rows_purchase_orders(db, tenant_id, filters)
    else:  # bom
        rows = await _rows_bom(db, tenant_id, bom_id, filters, indented, include_sub_assemblies)

    money_cols = MONEY_COLUMNS[entity] & set(resolved_columns)
    if currency and money_cols:
        rate = await _get_conversion_rate(db, tenant_id, currency)
        if rate != 1.0:
            for row in rows:
                for col in money_cols:
                    if row.get(col) is not None:
                        row[col] = round(row[col] * rate, 4)

    labels_by_key = {k: label for k, label, _ in COLUMNS[entity]}
    column_pairs = [(k, labels_by_key.get(k, k)) for k in resolved_columns]

    if format == "csv":
        content = _render_csv(column_pairs, rows)
    elif format == "json":
        content = _render_json(column_pairs, rows)
    elif format == "xlsx":
        content = _render_xlsx(column_pairs, rows, entity.replace("_", " ").title())
    else:  # pdf
        content = _render_pdf(column_pairs, rows, f"{entity.replace('_', ' ').title()} Export")

    ext = format
    stamp = date.today().strftime("%Y%m%d")
    id_part = f"-{bom_id}" if bom_id is not None else ""
    filename = f"{entity}{id_part}-{stamp}.{ext}"
    return content, CONTENT_TYPES[format], filename


def get_export_columns(entity: str) -> list[dict]:
    validate_entity(entity)
    return [{"key": k, "label": label, "default": default} for k, label, default in COLUMNS[entity]]


# ---- Templates ----


async def list_templates(db, tenant_id: Optional[int], entity: str) -> list[ExportTemplate]:
    validate_entity(entity)
    stmt = select(ExportTemplate).where(ExportTemplate.entity == entity)
    if tenant_id is not None:
        stmt = stmt.where(ExportTemplate.tenantId == tenant_id)
    stmt = stmt.order_by(ExportTemplate.name)
    return (await db.execute(stmt)).scalars().all()


async def get_template_or_404(db, tenant_id: Optional[int], template_id: int) -> ExportTemplate:
    stmt = select(ExportTemplate).where(ExportTemplate.id == template_id)
    if tenant_id is not None:
        stmt = stmt.where(ExportTemplate.tenantId == tenant_id)
    tmpl = (await db.execute(stmt)).scalar_one_or_none()
    if tmpl is None:
        raise HTTPException(status_code=404, detail="Export template not found")
    return tmpl


async def create_template(
    db, tenant_id: Optional[int], name: str, entity: str, config: dict
) -> ExportTemplate:
    validate_entity(entity)
    if "columns" in config and config["columns"] is not None:
        resolve_columns(entity, config["columns"])
    if "filters" in config and config["filters"]:
        validate_filters(entity, config["filters"])
    tmpl = ExportTemplate(name=name, entity=entity, config=config, tenantId=tenant_id)
    db.add(tmpl)
    await db.commit()
    await db.refresh(tmpl)
    return tmpl


async def delete_template(db, tenant_id: Optional[int], template_id: int) -> None:
    tmpl = await get_template_or_404(db, tenant_id, template_id)
    await db.delete(tmpl)
    await db.commit()
